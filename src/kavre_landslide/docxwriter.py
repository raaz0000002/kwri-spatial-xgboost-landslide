"""Minimal MDPI-style .docx writer (python-docx). Blocks are tuples produced by the manuscript builder."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Palatino Linotype'


def _set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge, val in kwargs.items():
        el = borders.find(qn(f'w:{edge}'))
        if el is None:
            el = OxmlElement(f'w:{edge}'); borders.append(el)
        for k, v in val.items():
            el.set(qn(f'w:{k}'), str(v))


def _runs(p, text, size=10, base_bold=False, base_italic=False, color=None, base_hl=False):
    """Inline markup: **bold**, *italic*, ^sup^, ~sub~, {{HL:text}} highlighted placeholder."""
    tokens = re.split(r'(\*\*.+?\*\*|\*[^*]+?\*|\^[^^]+?\^|~[^~]+?~|\{\{HL:.+?\}\})', text)
    for t in tokens:
        if not t:
            continue
        bold, ital, sup, sub, hl = base_bold, base_italic, False, False, base_hl
        if t.startswith('{{HL:'):
            _runs(p, t[5:-2], size, base_bold, base_italic, color, True); continue
        if t.startswith('**') and t.endswith('**'):
            t = t[2:-2]; bold = True
        elif t.startswith('*') and t.endswith('*') and len(t) > 2:
            t = t[1:-1]; ital = True
        elif t.startswith('^') and t.endswith('^'):
            t = t[1:-1]; sup = True
        elif t.startswith('~') and t.endswith('~'):
            t = t[1:-1]; sub = True
        r = p.add_run(t)
        r.font.name = FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        r.font.size = Pt(size); r.bold = bold; r.italic = ital
        r.font.superscript = sup; r.font.subscript = sub
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if hl:
            from docx.enum.text import WD_COLOR_INDEX
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def _small(p, maxw=2200):
    import os
    from PIL import Image
    os.makedirs('figs_doc', exist_ok=True); q = os.path.join('figs_doc', os.path.basename(p))
    im = Image.open(p).convert('RGB')
    if im.width > maxw: im = im.resize((maxw, int(im.height * maxw / im.width)), Image.LANCZOS)
    im.save(q, optimize=True); return q


def write(blocks, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    pf = st.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.15
    # line numbers (MDPI submissions are line-numbered)
    sectPr = sec._sectPr; ln = OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous'); sectPr.append(ln)
    for b in blocks:
        kind = b[0]
        if kind == 'title':
            p = doc.add_paragraph(); _runs(p, b[1], size=18, base_bold=True); p.paragraph_format.space_after = Pt(12)
        elif kind == 'authors':
            p = doc.add_paragraph(); _runs(p, b[1], size=10, base_bold=True); p.paragraph_format.space_after = Pt(6)
        elif kind == 'small':
            p = doc.add_paragraph(); _runs(p, b[1], size=8.5); p.paragraph_format.space_after = Pt(2)
        elif kind in ('h1', 'h2', 'h3'):
            lvl = {'h1': 1, 'h2': 2, 'h3': 3}[kind]
            p = doc.add_heading(level=lvl)
            _runs(p, b[1], size={1: 11, 2: 10, 3: 10}[lvl], base_bold=lvl < 3, base_italic=lvl >= 2, color='000000')
            p.paragraph_format.space_before = Pt(10 if lvl == 1 else 6); p.paragraph_format.space_after = Pt(4)
        elif kind == 'p':
            p = doc.add_paragraph(); _runs(p, b[1]); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(0.75) if (len(b) < 3 or b[2]) else None
            p.paragraph_format.space_after = Pt(2)
        elif kind == 'abstract':
            p = doc.add_paragraph(); _runs(p, '**Abstract:** ' + b[1], size=9); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
        elif kind == 'bullet':
            p = doc.add_paragraph(style='List Bullet'); _runs(p, b[1], size=10); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == 'caption':
            p = doc.add_paragraph(); _runs(p, b[1], size=9); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(4)
        elif kind == 'fig':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(_small(b[1]), width=Cm(b[3] if len(b) > 3 else 16.0))
            c = doc.add_paragraph(); _runs(c, b[2], size=9); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            c.paragraph_format.space_after = Pt(8)
        elif kind == 'table':
            caption, header, rows = b[1], b[2], b[3]
            notes = b[4] if len(b) > 4 else None
            c = doc.add_paragraph(); _runs(c, caption, size=9); c.paragraph_format.space_before = Pt(6); c.paragraph_format.space_after = Pt(3)
            t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]; cell.text = ''
                _runs(cell.paragraphs[0], str(h), size=8, base_bold=True)
                _set_cell_border(cell, top={'val': 'single', 'sz': 8}, bottom={'val': 'single', 'sz': 4})
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    cell = t.rows[i + 1].cells[j]; cell.text = ''
                    _runs(cell.paragraphs[0], str(v), size=8)
                    if i == len(rows) - 1:
                        _set_cell_border(cell, bottom={'val': 'single', 'sz': 8})
            if notes:
                n = doc.add_paragraph(); _runs(n, notes, size=8); n.paragraph_format.space_after = Pt(8)
            else:
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif kind == 'ref':
            p = doc.add_paragraph(); _runs(p, b[1], size=9); p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6); p.paragraph_format.space_after = Pt(1)
        elif kind == 'pagebreak':
            doc.add_page_break()
    doc.save(path)
